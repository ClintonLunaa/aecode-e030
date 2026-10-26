"""Reportes PDF, Excel y CSV del análisis ejecutado."""
from dataclasses import asdict
from io import BytesIO
from html import escape
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from core.contratos import VERSION, UNIDADES, FUENTES, HISTORICO, INVALIDO
from diagnosticos import resumen_calculos, estado_claro, nombre_edicion, notas_de_revision

def tablas_exportables(r):
    """Nombres de columnas no inducen a amplificación reglamentaria."""
    tablas={**r.tablas,"Resumen_Modal":r.tabla_modal}
    tablas["Combinacion"]=r.combinacion.rename(columns={
        "Vreal_kN":"Referencia_ET1_no_diseno_kN",
        "Vrnc_h_kN":"Vcombinado_25_75_kN","Vrcsc_kN":"V_SRSS_kN",
    })
    return tablas

def _permitir(r):
    if r.estado == INVALIDO:
        raise ValueError("Escenario INVALIDO: exporte solo el diagnóstico, no reportes de resultados.")

def csv_contextual(r,nombre):
    _permitir(r)
    df=tablas_exportables(r)[nombre].copy()
    df["Tipo_de_analisis"]="Respuesta modal con modos ingresados";df["Edicion"]=nombre_edicion(r.edicion);df["Referencia"]=r.firma[:12]
    return df.to_csv(index=False).encode("utf-8-sig")

def _figuras_estaticas(r):
    """Figuras legibles tanto en el PDF como en la hoja Excel."""
    import numpy as np
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from core.espectros import ordenada_g

    azul, violeta, ocre = "#12718C", "#6E5A9C", "#C77830"
    with plt.rc_context({
        "font.family": "DejaVu Sans", "font.size": 10,
        "axes.labelsize": 10.5, "xtick.labelsize": 9.5,
        "ytick.labelsize": 9.5, "axes.edgecolor": "#BACAD5",
        "axes.labelcolor": "#334C60", "text.color": "#2B4052",
        "xtick.color": "#51697B", "ytick.color": "#51697B",
        "axes.spines.top": False, "axes.spines.right": False,
        "figure.facecolor": "white", "axes.facecolor": "white",
    }):
        figuras = []
        tiempos = np.linspace(.0001, max(3.2, r.parametros["tl"] * 1.1), 500)
        valores = ordenada_g(tiempos, r.parametros, r.edicion)
        fig, ax = plt.subplots(figsize=(8.4, 3.4))
        ax.plot(tiempos, valores, color=azul, linewidth=2.6)
        marcas = (azul, ocre, violeta)
        for i, periodo in enumerate(r.periodos):
            ax.scatter(float(periodo), float(r.espectro.Sa_m_s2.iloc[i]/r.parametros["gravedad"]),
                       s=51, c=marcas[i], edgecolor="white", linewidth=1.3, zorder=4,
                       label=f"Modo {i+1}: {periodo:.3f} s")
        ax.set(xlabel="Período, T (s)", ylabel="Sa/g")
        ax.set_xlim(0, max(tiempos)); ax.set_ylim(bottom=0)
        ax.grid(alpha=.16, color="#7890A3")
        ax.legend(loc="upper right", frameon=False, fontsize=8.8, ncol=1,
                  labelspacing=.35, handlelength=.8)
        fig.subplots_adjust(left=.09,right=.98,top=.97,bottom=.18)
        b = BytesIO(); fig.savefig(b, format="png", dpi=165, facecolor="white")
        plt.close(fig); b.seek(0); figuras.append(b)

        fig, ax = plt.subplots(figsize=(8.4, 3.4))
        for j, color in enumerate(marcas):
            ax.barh(r.datos.Piso + (j-1)*.23, r.fuerzas[:, j], height=.21,
                    label=f"Modo {j+1}", color=color)
        ax.axvline(0, color="#536A7B", linewidth=.9)
        ax.set(xlabel="Fuerza lateral por modo (kN)", ylabel="Piso")
        ax.set_yticks(list(r.datos.Piso))
        ax.grid(axis="x",alpha=.15,color="#7890A3")
        ax.set_axisbelow(True)
        ax.legend(loc="upper center",bbox_to_anchor=(.52,1.13),ncol=3,
                  fontsize=9,frameon=False,columnspacing=1.5)
        fig.subplots_adjust(left=.09,right=.98,top=.83,bottom=.18)
        b = BytesIO(); fig.savefig(b, format="png", dpi=165, facecolor="white")
        plt.close(fig); b.seek(0); figuras.append(b)
        return figuras

def crear_excel(r):
    _permitir(r)
    wb=Workbook();s=wb.active;s.title="Resumen"
    k=resumen_calculos(r)
    filas=[
        ["AECODE - Análisis modal espectral"],
        ["Tipo de análisis","Respuesta modal con datos ingresados"],["Edición",nombre_edicion(r.edicion)],
        ["Fecha",r.fecha_utc[:10]],["Origen de los datos",r.procedencia],
        ["Vcombinado 25/75 (kN)",k["v_combinada"]],
        ["SRSS basal (kN)",k["v_srss"]],["Masa total (t)",k["masa_total_t"]],
        ["Masa efectiva (%)",k["acumulada_pct"]],["T1 ingresado (s)",r.periodos[0]],
        [HISTORICO,k["v_et1"]],["Fórmula histórica","0.75 R Vcombinado; NO DISEÑO"],
        ["Cortante mínimo",notas_de_revision(r)[5][1]],
    ]
    for row in filas:s.append(row)
    for name,tabla in tablas_exportables(r).items():
        sh=wb.create_sheet(name[:31])
        sh.append(["Respuesta modal con datos ingresados",nombre_edicion(r.edicion),"Referencia: " + r.firma[:12]]);sh.append(list(tabla.columns))
        for row in tabla.itertuples(index=False,name=None):
            sh.append([x.item() if hasattr(x,"item") else x for x in row])
    sh=wb.create_sheet("Observaciones")
    sh.append(["Tema","Interpretación"])
    for tema, comentario in notas_de_revision(r):sh.append([tema,comentario])
    sh=wb.create_sheet("Escenario")
    sh.append(["Campo","Valor","Unidad/contexto"])
    for key,val in r.parametros.items():sh.append([key,val,"Parámetro declarado"])
    for key,val in {"edicion":nombre_edicion(r.edicion),"procedencia":r.procedencia,
                    "referencia":r.firma[:12]}.items():
        sh.append([key,val,""])
    for key,val in UNIDADES.items():sh.append(["Unidad "+key,val,"Declarada"])
    premisas = r.premisas_minimo or {}
    sh.append(["Premisas mínimo", "; ".join(f"{clave}: {valor}" for clave, valor in premisas.items()), ""])
    sh=wb.create_sheet("Graficos")
    sh.append(["Gráficos del mismo escenario",r.edicion,r.firma])
    imgs=_figuras_estaticas(r)
    for anchor,buf in zip(["A4","A27"],imgs):
        im=XLImage(buf);im.width=720;im.height=333;sh.add_image(im,anchor)
    for sh in wb:
        sh.freeze_panes="A3"
        for c in sh[1]:
            c.fill=PatternFill("solid",fgColor="163B4D")
            c.font=Font(color="FFFFFF",bold=True)
        for col in sh.columns:
            letra=col[0].column_letter
            sh.column_dimensions[letra].width=38 if letra=="A" else 30
        if sh.title in ("Resumen","Observaciones","Escenario"):
            sh.column_dimensions["B"].width=55
            sh.column_dimensions["D"].width=100
            for row in sh:
                for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
                sh.row_dimensions[row[0].row].height=48
        sh.sheet_view.showGridLines=False
        sh.page_setup.orientation="landscape";sh.page_setup.paperSize=sh.PAPERSIZE_A4
        sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
        sh.sheet_properties.pageSetUpPr.fitToPage=True
        if sh.title!="Graficos":sh.print_area=sh.dimensions
    wb["Graficos"].print_area="A1:J49"
    out=BytesIO();wb.save(out);return out.getvalue()

def crear_pdf(r, capturas=()):
    """Informe académico del escenario activo, en formato A4 de tres páginas."""
    _permitir(r)
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, PageBreak, Image, KeepTogether, HRFlowable)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from matplotlib import get_data_path
    from PIL import Image as PILImage

    carpeta_fuentes = Path(get_data_path()) / "fonts" / "ttf"
    if "AEC-Regular" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AEC-Regular", str(carpeta_fuentes / "DejaVuSans.ttf")))
    if "AEC-Semibold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("AEC-Semibold", str(carpeta_fuentes / "DejaVuSans-Bold.ttf")))

    azul = colors.HexColor("#153452")
    azul_medio = colors.HexColor("#196B95")
    texto = colors.HexColor("#344A5D")
    gris = colors.HexColor("#617384")
    linea = colors.HexColor("#DEE6EC")
    suave = colors.HexColor("#F3F7FA")
    blanco = colors.white

    estilos = getSampleStyleSheet()
    estilos.add(ParagraphStyle(name="EyebrowAE", fontName="AEC-Semibold", fontSize=8,
                               leading=12, textColor=azul_medio, spaceAfter=6))
    estilos.add(ParagraphStyle(name="TitAE", fontName="AEC-Semibold", fontSize=21.5,
                               leading=28, textColor=azul, spaceAfter=7))
    estilos.add(ParagraphStyle(name="SubAE", fontName="AEC-Regular", fontSize=10,
                               leading=16, textColor=gris, spaceAfter=10))
    estilos.add(ParagraphStyle(name="H2AE", fontName="AEC-Semibold", fontSize=10.5,
                               leading=15, textColor=azul, spaceBefore=14, spaceAfter=7))
    estilos.add(ParagraphStyle(name="H3AE", fontName="AEC-Semibold", fontSize=9.2,
                               leading=14, textColor=azul, spaceAfter=4))
    estilos.add(ParagraphStyle(name="BodyAE", fontName="AEC-Regular", fontSize=8.75,
                               leading=14.5, textColor=texto, spaceAfter=8))
    estilos.add(ParagraphStyle(name="SmallAE", fontName="AEC-Regular", fontSize=7.65,
                               leading=12.5, textColor=gris, spaceAfter=5))
    estilos.add(ParagraphStyle(name="CellAE", fontName="AEC-Regular", fontSize=8,
                               leading=12, textColor=texto))
    estilos.add(ParagraphStyle(name="CellHeadAE", fontName="AEC-Semibold", fontSize=7.8,
                               leading=12, textColor=azul))
    estilos.add(ParagraphStyle(name="ValueAE", fontName="AEC-Semibold", fontSize=18,
                               leading=25, textColor=azul))
    estilos.add(ParagraphStyle(name="LabelAE", fontName="AEC-Semibold", fontSize=7.2,
                               leading=12, textColor=azul_medio))

    def texto_pdf(valor, estilo="BodyAE"):
        return Paragraph(escape(str(valor)), estilos[estilo])

    def tabla(filas, anchos, alineacion=None):
        celdas = [
            [texto_pdf(valor, "CellHeadAE" if i == 0 else "CellAE") for valor in fila]
            for i, fila in enumerate(filas)
        ]
        t = Table(celdas, colWidths=anchos, repeatRows=1, hAlign="LEFT")
        reglas = [
            ("BACKGROUND", (0,0), (-1,0), suave),
            ("LINEBELOW", (0,0), (-1,0), 0.8, azul_medio),
            ("LINEBELOW", (0,1), (-1,-1), 0.35, linea),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("LEFTPADDING", (0,0), (-1,-1), 9),
            ("RIGHTPADDING", (0,0), (-1,-1), 8),
            ("TOPPADDING", (0,0), (-1,-1), 7.5),
            ("BOTTOMPADDING", (0,0), (-1,-1), 7.5),
        ]
        for i in range(2, len(filas), 2):
            reglas.append(("BACKGROUND", (0,i), (-1,i), colors.HexColor("#FBFCFD")))
        if alineacion:
            reglas.append(("ALIGN",(alineacion,1),(alineacion,-1),"RIGHT"))
        t.setStyle(TableStyle(reglas))
        return t

    def cabecera(kicker, titulo, subtitulo=None):
        contenido.append(texto_pdf(kicker, "EyebrowAE"))
        contenido.append(texto_pdf(titulo, "TitAE"))
        if subtitulo:
            contenido.append(texto_pdf(subtitulo, "SubAE"))
        contenido.append(HRFlowable(width="100%", thickness=0.7, color=linea, spaceBefore=2, spaceAfter=12))

    def figura_ajustada(buffer, ancho_max=503, alto_max=206):
        buffer.seek(0)
        with PILImage.open(buffer) as img:
            w, h = img.size
        factor = min(ancho_max / w, alto_max / h)
        return Image(buffer, width=w * factor, height=h * factor, hAlign="CENTER")

    k = resumen_calculos(r)
    ancho = 503
    contenido = []
    edicion = nombre_edicion(r.edicion)

    # PÁGINA 1: resumen claro del trabajo y del escenario utilizado.
    cabecera("AECODE  /  PYTHON APLICADO A INGENIERÍA CIVIL", "Análisis modal espectral",
             "Entregable 2  ·  Modelo académico de " + str(len(r.datos)) + " pisos  ·  " + edicion)
    contenido.append(texto_pdf("Clinton Luna Mamani Vargas  |  Universidad de Piura", "SmallAE"))
    contenido.append(Spacer(1, 7))
    contenido.append(texto_pdf(
        "Se calcula la respuesta sísmica de un modelo de cinco pisos a partir de las masas, "
        "los períodos y las formas modales proporcionados en el ejercicio. Los resultados "
        "incluyen el espectro de aceleraciones, la participación modal, las fuerzas "
        "laterales y los cortantes de entrepiso."
    ))
    indicadores = [
        ("CORTANTE COMBINADO", f'{k["v_combinada"]:.2f} kN', "Combinación modal 25/75"),
        ("MASA EFECTIVA", f'{k["acumulada_pct"]:.2f} %', "Participación acumulada"),
        ("PRIMER PERÍODO", f'{r.periodos[0]:.3f} s', "Dato suministrado"),
    ]
    cajas = []
    for etiqueta, valor, nota in indicadores:
        cajas.append(Table([
            [texto_pdf(etiqueta, "LabelAE")],
            [texto_pdf(valor, "ValueAE")],
            [texto_pdf(nota, "SmallAE")],
        ], colWidths=[131], style=TableStyle([
            ("VALIGN", (0,0),(-1,-1),"TOP"),
            ("LEFTPADDING",(0,0),(-1,-1),0),
            ("RIGHTPADDING",(0,0),(-1,-1),0),
            ("TOPPADDING",(0,0),(-1,-1),1),
            ("BOTTOMPADDING",(0,0),(-1,-1),1),
        ])))
    resumen = Table([cajas], colWidths=[ancho/3]*3, hAlign="LEFT")
    resumen.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),suave),
        ("LINEBELOW",(0,0),(-1,0),1.2,azul_medio),
        ("INNERGRID",(0,0),(-1,-1),0.4,linea),
        ("LEFTPADDING",(0,0),(-1,-1),13),
        ("RIGHTPADDING",(0,0),(-1,-1),6),
        ("TOPPADDING",(0,0),(-1,-1),12),
        ("BOTTOMPADDING",(0,0),(-1,-1),11),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
    ]))
    contenido.append(Spacer(1,9)); contenido.append(resumen)

    contenido.append(texto_pdf("Datos de entrada", "H2AE"))
    contenido.append(tabla([
        ["Característica", "Valor", "Parámetro sísmico", "Valor"],
        ["Número de pisos", str(len(r.datos)), "Factor de zona (Z)", f'{r.parametros["z"]:g}'],
        ["Modos considerados", str(len(r.periodos)), "Factor de uso (U)", f'{r.parametros["u"]:g}'],
        ["Masa total",f'{k["masa_total_t"]:.2f} t',"Factor de suelo (S)",f'{r.parametros["s"]:g}'],
        ["Altura de entrepiso",f'{r.parametros["altura_piso"]:.2f} m',"Reducción (R)",f'{r.parametros["r"]:g}'],
        ["Período Tp",f'{r.parametros["tp"]:g} s',"Período TL",f'{r.parametros["tl"]:g} s'],
    ], [135, 108, 151, 109]))

    contenido.append(texto_pdf("Resultados principales", "H2AE"))
    contenido.append(tabla([
        ["Magnitud", "Resultado"],
        ["Cortante basal por SRSS",f'{k["v_srss"]:.2f} kN'],
        ["Cortante basal combinado (25/75)",f'{k["v_combinada"]:.2f} kN'],
        ["Masa efectiva acumulada",f'{k["acumulada_pct"]:.2f} %'],
    ],[337,166]))
    contenido.append(Spacer(1,10))
    contenido.append(texto_pdf(
        "Los períodos y las formas modales son datos de entrada. La aplicación no realiza "
        "la extracción modal a partir de una matriz de rigidez.", "SmallAE"))
    contenido.append(PageBreak())

    # PÁGINA 2: procedimiento y tablas de resultados calculados.
    cabecera("PROCEDIMIENTO DE CÁLCULO", "Del espectro a los cortantes",
             "Síntesis de las operaciones utilizadas en la aplicación")
    pasos = [
        ("01", "Espectro de aceleraciones",
         "Para cada período se evalúa Sa(T) = Z·U·S·C(T)·g/R, "
         "con el coeficiente C(T) de la edición seleccionada."),
        ("02", "Participación modal",
         "Con las masas y las formas de entrada se obtiene el factor "
         "Γ = (φᵀM1)/(φᵀMφ) y la masa efectiva asociada a cada modo."),
        ("03", "Fuerzas y cortantes",
         "La fuerza por nivel se calcula como F(i,r) = m(i)·φ(i,r)·Γ(r)·Sa(r). "
         "Los cortantes se obtienen acumulando las fuerzas desde la cubierta."),
        ("04", "Combinación modal",
         "Se aplica 0.25·Σ|Vr| + 0.75·√Σ(Vr²) a los cortantes "
         "de cada piso. La expresión conserva los aportes de todos los modos."),
    ]
    for num, nombre, descripcion in pasos:
        fila = Table([[
            texto_pdf(num, "EyebrowAE"),
            [texto_pdf(nombre, "H3AE"),texto_pdf(descripcion)],
        ]], colWidths=[39, 464], hAlign="LEFT")
        fila.setStyle(TableStyle([
            ("VALIGN",(0,0),(-1,-1),"TOP"),
            ("LINEBELOW",(0,0),(-1,-1),0.4,linea),
            ("TOPPADDING",(0,0),(-1,-1),8),
            ("BOTTOMPADDING",(0,0),(-1,-1),6),
            ("LEFTPADDING",(0,0),(-1,-1),0),
            ("RIGHTPADDING",(0,0),(-1,-1),6),
        ]))
        contenido.append(KeepTogether([fila,Spacer(1,2)]))

    contenido.append(texto_pdf("Resultados por modo", "H2AE"))
    contenido.append(tabla([
        ["Modo", "T (s)", "C(T)", "Sa (m/s²)", "Masa efectiva (%)"]
    ] + [[str(i+1), f'{r.periodos[i]:.3f}', f'{r.espectro.C.iloc[i]:.4f}',
          f'{r.espectro.Sa_m_s2.iloc[i]:.4f}',
          f'{r.participacion.Participacion_pct.iloc[i]:.2f}']
         for i in range(len(r.periodos))], [57,78,91,123,154]))
    contenido.append(texto_pdf("Cortantes por piso", "H2AE"))
    contenido.append(tabla([["Piso", "SRSS (kN)", "Combinación 25/75 (kN)"]] + [
        [str(i+1), f'{r.combinacion.Vrcsc_kN.iloc[i]:.3f}',
         f'{r.combinacion.Vrnc_h_kN.iloc[i]:.3f}']
        for i in range(len(r.datos))], [102,172,229]))
    contenido.append(Spacer(1,13))
    contenido.append(texto_pdf(
        "En este escenario no se aplica un ajuste por cortante mínimo. Para efectuar "
        "esa comparación deben indicarse y justificarse los parámetros adicionales "
        "en la sección Escenario.", "SmallAE") if r.minimo["estado"] == "NO_EVALUADO" else texto_pdf(
        "La comparación del cortante mínimo utiliza las hipótesis indicadas "
        "por el usuario en la sección Escenario. Las fuerzas no se escalan automáticamente.", "SmallAE"))
    contenido.append(PageBreak())

    # PÁGINA 3: gráficos sin deformar e interpretación académica.
    cabecera("RESULTADOS GRÁFICOS", "Lectura de la respuesta modal",
             "Espectro reducido y fuerzas laterales del escenario seleccionado")
    imagenes = _figuras_estaticas(r)
    for titulo, imagen, nota in [
        ("Espectro de aceleraciones reducido", imagenes[0],
         "Cada marcador corresponde al período de uno de los modos ingresados."),
        ("Fuerzas laterales por modo", imagenes[1],
         "Las fuerzas mantienen el signo de la respuesta de cada modo."),
    ]:
        contenido.append(texto_pdf(titulo,"H2AE"))
        contenido.append(figura_ajustada(imagen, ancho_max=ancho, alto_max=210))
        contenido.append(Spacer(1,4))
        contenido.append(texto_pdf(nota, "SmallAE"))

    contenido.append(texto_pdf("Interpretación", "H2AE"))
    contenido.append(texto_pdf(
        f'En el caso analizado, los tres modos representan el {k["acumulada_pct"]:.2f} % '
        "de la masa ingresada. Este valor describe la participación de los modos "
        "suministrados; para determinar los modos desde el edificio se necesitarían "
        "sus propiedades de rigidez, masa y apoyos. Las verificaciones de diseño "
        "de los elementos estructurales no forman parte del ejercicio."))
    contenido.append(texto_pdf("Referencia normativa: " + FUENTES[r.edicion], "SmallAE"))

    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida, pagesize=(595.28, 841.89), leftMargin=46, rightMargin=46,
        topMargin=48, bottomMargin=49,
        title="Análisis modal espectral - Entregable 2",
        author="Clinton Luna Mamani Vargas")

    def decoracion(canvas, doc):
        canvas.saveState()
        ancho_pagina, alto_pagina = (595.28, 841.89)
        canvas.setFillColor(azul)
        canvas.rect(0, alto_pagina-9, ancho_pagina, 9, stroke=0, fill=1)
        canvas.setStrokeColor(linea)
        canvas.setLineWidth(0.6)
        canvas.line(46,37,549.28,37)
        canvas.setFont("AEC-Regular",7.3)
        canvas.setFillColor(gris)
        canvas.drawString(46,25,"AECODE  |  Análisis modal espectral")
        canvas.drawRightString(549.28,25,f"{doc.page} / 3")
        canvas.restoreState()

    documento.build(contenido, onFirstPage=decoracion, onLaterPages=decoracion)
    return salida.getvalue()
